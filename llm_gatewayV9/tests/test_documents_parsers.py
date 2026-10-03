"""Parser tests, with fixtures generated in-process.

Fixtures are built here rather than committed as binaries so the expectations
stay readable and the suite has no opaque files. Each parser's contract:
preserve STRUCTURE (headings, lists, tables, page numbers) and report what it
could not read instead of failing silently.
"""
from __future__ import annotations

import csv
import io
import sys
import zipfile
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))

from documents.chunker import Block, chunk_blocks  # noqa: E402
from documents import parsers as P  # noqa: E402

MD = """# Title

Intro paragraph with two sentences. It continues here.

## Method

- first bullet item
- second bullet item

Some prose after the list. More prose follows it.

```python
def f():
    return 1
```

| Name | Value |
|------|------:|
| alpha | 1 |
| beta | 2 |

## Results

Final paragraph of the document. Nothing more to add.
"""


# ── dispatch ─────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("name, expect", [
    ("a.pdf", "pdf"), ("a.docx", "docx"), ("a.md", "md"), ("a.txt", "txt"),
    ("a.html", "html"), ("a.csv", "csv"), ("a.tsv", "csv"), ("a.xlsx", "xlsx"),
])
def test_extension_dispatch(name, expect):
    assert P.detect_type(name) == expect


@pytest.mark.parametrize("name", ["a.exe", "a.zip", "a", "a.pptx", "a.doc"])
def test_unsupported_types_are_rejected_loudly(name):
    with pytest.raises(P.UnsupportedDocument) as e:
        P.detect_type(name)
    # The message must name what IS supported, not just refuse.
    assert "PDF" in str(e.value)


# ── markdown ─────────────────────────────────────────────────────────────────

def test_markdown_preserves_structure():
    res = P.parse("doc.md", MD.encode())
    kinds = [b.kind for b in res.blocks]
    assert "heading" in kinds and "list" in kinds
    assert "table" in kinds and "code" in kinds
    assert not res.warnings, res.warnings


def test_markdown_heading_levels():
    res = P.parse("doc.md", MD.encode())
    headings = {b.text: b.level for b in res.blocks if b.kind == "heading"}
    assert headings["Title"] == 1
    assert headings["Method"] == 2
    assert headings["Results"] == 2


def test_markdown_list_is_kept_whole():
    res = P.parse("doc.md", MD.encode())
    lists = [b for b in res.blocks if b.kind == "list"]
    assert lists and "first bullet item" in lists[0].text
    assert "second bullet item" in lists[0].text


def test_markdown_table_header_and_rows():
    res = P.parse("doc.md", MD.encode())
    t = next(b for b in res.blocks if b.kind == "table")
    assert t.header == ["Name", "Value"]
    assert ["alpha", "1"] in t.rows
    assert ["beta", "2"] in t.rows


def test_markdown_code_fence_is_not_reinterpreted():
    """A '#' inside a fence is code, not a heading."""
    src = "# Real\n\n```\n# not a heading\nstill code\n```\n\nAfter.\n"
    res = P.parse("d.md", src.encode())
    headings = [b.text for b in res.blocks if b.kind == "heading"]
    assert headings == ["Real"]


def test_markdown_setext_headings():
    res = P.parse("d.md", b"Chapter One\n==========\n\nBody text here.\n")
    h = [b for b in res.blocks if b.kind == "heading"]
    assert h and h[0].text == "Chapter One" and h[0].level == 1


# ── plain text ───────────────────────────────────────────────────────────────

def test_text_detects_headings():
    body = ("INTRODUCTION\n\nFirst paragraph. Second sentence here.\n\n"
            "1.2 Some Section\n\nMore body text. It continues.\n")
    res = P.parse("a.txt", body.encode())
    heads = {b.text for b in res.blocks if b.kind == "heading"}
    assert "INTRODUCTION" in heads
    assert "1.2 Some Section" in heads


def test_text_paragraphs_are_not_headings():
    body = ("This is an ordinary sentence that happens to be long enough to "
            "look like prose. It should stay a paragraph.\n")
    res = P.parse("a.txt", body.encode())
    assert res.blocks[0].kind == "para"


def test_empty_file_reports_no_text():
    res = P.parse("a.txt", b"   \n\n  ")
    assert not res.blocks
    assert any("no readable text" in w for w in res.warnings)


def test_weird_whitespace_is_normalised():
    src = "Heading\r\n\r\nBody with\xa0non-breaking and soft\u00adhyphen.\r\n"
    res = P.parse("a.txt", src.encode())
    text = " ".join(b.text for b in res.blocks)
    assert "\xa0" not in text
    assert "\u00ad" not in text
    assert "non-breaking" in text and "soft" in text


# ── HTML ─────────────────────────────────────────────────────────────────────

HTML = """<html><head><style>.x{}</style><script>evil()</script></head>
<body>
<nav>Home About Contact</nav>
<main>
  <h1>Page Title</h1>
  <p>First paragraph. Second sentence.</p>
  <h2>Section</h2>
  <ul><li>alpha</li><li>beta</li></ul>
  <table><tr><th>Col</th><th>Val</th></tr><tr><td>x</td><td>1</td></tr></table>
</main>
<footer>Copyright</footer>
</body></html>"""


def test_html_keeps_headings_and_drops_chrome():
    res = P.parse("a.html", HTML.encode())
    heads = {b.text: b.level for b in res.blocks if b.kind == "heading"}
    assert heads.get("Page Title") == 1
    assert heads.get("Section") == 2
    text = " ".join(b.text for b in res.blocks)
    assert "Home About Contact" not in text, "nav leaked into content"
    assert "Copyright" not in text, "footer leaked into content"
    assert "evil()" not in text, "script leaked into content"


def test_html_list_items_are_blocks():
    res = P.parse("a.html", HTML.encode())
    assert any(b.kind == "list" for b in res.blocks)


def test_html_table_becomes_a_table_block():
    res = P.parse("a.html", HTML.encode())
    t = next((b for b in res.blocks if b.kind == "table"), None)
    assert t is not None and t.header == ["Col", "Val"]


# ── CSV ──────────────────────────────────────────────────────────────────────

def test_csv_header_and_rows():
    src = "region,docs,gb\nnorth,10,3.2\nsouth,20,4.4\neast,30,9.1\n"
    res = P.parse("a.csv", src.encode())
    t = res.blocks[0]
    assert t.kind == "table"
    assert t.header == ["region", "docs", "gb"]
    assert ["north", "10", "3.2"] in t.rows
    assert res.meta["cols"] == 3


def test_tsv_is_sniffed():
    res = P.parse("a.tsv", b"a\tb\tc\n1\t2\t3\n")
    assert res.blocks[0].header == ["a", "b", "c"]


def test_semicolon_csv():
    res = P.parse("a.csv", b"a;b;c\n1;2;3\n")
    assert res.blocks[0].header == ["a", "b", "c"]


def test_ragged_rows_are_padded_not_dropped():
    src = "a,b,c\n1,2\n3,4,5,6\n"
    res = P.parse("a.csv", src.encode())
    t = res.blocks[0]
    assert all(len(r) == 4 for r in t.rows), t.rows


def test_csv_header_is_repeated_on_every_chunk():
    """The point of a header: a chunk of numbers alone is meaningless."""
    src = "region,units,revenue\n" + "".join(
        f"r{i},{i * 3},{i * 100}\n" for i in range(30))
    res = P.parse("a.csv", src.encode())
    chunks = [c for c in chunk_blocks(res.blocks, target_words=30)
              if c.kind == "table"]
    assert len(chunks) > 1
    for c in chunks:
        assert "region" in c.text and "revenue" in c.text


# ── PDF ──────────────────────────────────────────────────────────────────────

def _make_pdf(pages: list[list[tuple[str, float]]]) -> bytes:
    """Minimal but VALID PDF with real text operators, so the parser's
    font-size heuristic has something to work with.

    Object numbering matters here: the catalog and page tree must be
    referenced correctly or pypdf returns a document whose pages extract as
    empty strings -- which looks exactly like a PDF parsing failure.
    """
    import zlib

    # 1 catalog, 2 page tree, then per page: page obj + contents obj,
    # then the font at the end.
    n_pages = len(pages)
    font_no = 3 + 2 * n_pages
    objs: dict[int, bytes] = {}

    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_pages))
    objs[2] = (f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>"
               .encode())
    for i, lines in enumerate(pages):
        page_no, cont_no = 3 + 2 * i, 4 + 2 * i
        parts = ["BT"]
        y = 720
        for text, size in lines:
            esc = (text.replace("\\", r"\\").replace("(", r"\(")
                       .replace(")", r"\)"))
            parts.append(f"/F1 {size} Tf 1 0 0 1 56 {y} Tm ({esc}) Tj")
            y -= int(size) + 8
        parts.append("ET")
        stream = zlib.compress("\n".join(parts).encode("latin-1", "replace"))
        objs[page_no] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_no} 0 R >> >> "
            f"/Contents {cont_no} 0 R >>".encode())
        objs[cont_no] = (b"<< /Length " + str(len(stream)).encode() +
                         b" /Filter /FlateDecode >>\nstream\n" + stream +
                         b"\nendstream")
    objs[font_no] = (b"<< /Type /Font /Subtype /Type1 "
                     b"/BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for n in sorted(objs):
        offsets[n] = out.tell()
        out.write(f"{n} 0 obj\n".encode() + objs[n] + b"\nendobj\n")
    size = max(objs) + 1
    xref = out.tell()
    out.write(f"xref\n0 {size}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for n in range(1, size):
        out.write(f"{offsets.get(n, 0):010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {size} /Root 1 0 R >>\n"
              f"startxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_pdf_extracts_text_with_page_numbers():
    pdf = _make_pdf([
        [("Chapter One", 22.0), ("Body text on page one. More here.", 11.0)],
        [("Chapter Two", 22.0), ("Body text on page two. And more.", 11.0)],
    ])
    res = P.parse("a.pdf", pdf)
    text = " ".join(b.text for b in res.blocks)
    assert "Chapter One" in text and "Chapter Two" in text
    assert res.meta.get("pages") == 2
    pages = {b.page for b in res.blocks if b.page}
    assert pages == {1, 2}, pages


def test_pdf_large_font_becomes_a_heading():
    pdf = _make_pdf([
        [("The Big Heading", 24.0), ("Ordinary sentence follows. And more.", 10.0)],
    ])
    res = P.parse("a.pdf", pdf)
    heads = [b for b in res.blocks if b.kind == "heading"]
    assert any("Big Heading" in h.text for h in heads), \
        [(b.kind, b.text[:40]) for b in res.blocks]


def test_pdf_pages_are_reassembled_into_sentences():
    """A PDF splits a sentence across lines; the parser must rejoin it."""
    pdf = _make_pdf([
        [("The first part of a sentence that", 11.0),
         ("continues onto the next line. Done.", 11.0)],
    ])
    res = P.parse("a.pdf", pdf)
    text = " ".join(b.text for b in res.blocks)
    assert "continues onto the next line" in text
    # And it should read as one sentence, not two fragments.
    from documents.chunker import split_sentences
    sents = split_sentences(text)
    assert any("next line" in s for s in sents)


def test_corrupt_pdf_reports_instead_of_raising():
    res = P.parse("a.pdf", b"not a pdf at all")
    assert not res.blocks
    assert res.warnings, "a corrupt PDF must say so"


# ── DOCX ─────────────────────────────────────────────────────────────────────

def _make_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    """Build a .docx by hand: a zip with the minimum parts Word needs."""
    import xml.sax.saxutils as sx

    body = []
    for text, style in paragraphs:
        ppr = (f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>'
               if style else "")
        body.append(
            f'<w:p>{ppr}<w:r><w:t xml:space="preserve">'
            f'{sx.escape(text)}</w:t></w:r></w:p>')
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main"><w:body>'
        + "".join(body) +
        '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Col</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t>Val</w:t></w:r></w:p></w:tc></w:tr>'
        '<w:tr><w:tc><w:p><w:r><w:t>x</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t>1</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
        '</w:body></w:document>')
    # python-docx refuses the file unless the officeDocument relationship and
    # the matching content-type override are present, so both are declared.
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/'
          'content-types">'
          '<Default Extension="rels" ContentType="application/vnd.'
          'openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          f'<Override PartName="/word/document.xml" ContentType='
          f'"application/vnd.openxmlformats-officedocument.'
          f'wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/'
            'package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>')
    styles = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              f'<w:styles xmlns:w="{W}">'
              f'<w:style w:type="paragraph" w:styleId="Heading1">'
              f'<w:name w:val="heading 1"/></w:style>'
              f'<w:style w:type="paragraph" w:styleId="Heading2">'
              f'<w:name w:val="heading 2"/></w:style>'
              f'</w:styles>')
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)
        z.writestr("word/styles.xml", styles)
    return out.getvalue()


def test_docx_uses_paragraph_styles_as_headings():
    raw = _make_docx([
        ("Document Title", "Heading1"),
        ("Body sentence one. Body sentence two.", "Normal"),
        ("A Subsection", "Heading2"),
        ("More body text here. It ends.", "Normal"),
    ])
    res = P.parse("a.docx", raw)
    heads = {b.text: b.level for b in res.blocks if b.kind == "heading"}
    assert heads.get("Document Title") == 1
    assert heads.get("A Subsection") == 2


def test_docx_table_is_captured():
    raw = _make_docx([("Body.", "Normal")])
    res = P.parse("a.docx", raw)
    t = next((b for b in res.blocks if b.kind == "table"), None)
    assert t is not None and t.header == ["Col", "Val"]


def test_corrupt_docx_reports_instead_of_raising():
    res = P.parse("a.docx", b"PK\x03\x04 truncated garbage")
    assert not res.blocks
    assert res.warnings


# ── XLSX ─────────────────────────────────────────────────────────────────────

def test_xlsx_reads_sheets_when_openpyxl_present():
    openpyxl = pytest.importorskip("openpyxl")
    import openpyxl as _x
    wb = _x.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["region", "units"])
    ws.append(["north", 10])
    ws.append(["south", 20])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["note"])
    ws2.append(["hello"])
    buf = io.BytesIO()
    wb.save(buf)
    res = P.parse("a.xlsx", buf.getvalue())
    tables = [b for b in res.blocks if b.kind == "table"]
    assert len(tables) == 2, [b.meta for b in tables]
    assert tables[0].meta["sheet"] == "Sales"
    assert tables[0].header == ["region", "units"]
    assert ["north", "10"] in tables[0].rows


def _xlsx_bytes() -> bytes:
    openpyxl = pytest.importorskip("openpyxl")
    import openpyxl as _x
    wb = _x.Workbook()
    ws = wb.active
    ws.append(["col_a", "col_b"])
    ws.append(["1", "2"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_falls_back_without_openpyxl(monkeypatch):
    """No openpyxl must still produce content, and must say it is degraded."""
    import builtins
    real_import = builtins.__import__

    def fake(name, *a, **k):
        if name == "openpyxl":
            raise ImportError("blocked for the test")
        return real_import(name, *a, **k)

    raw = _xlsx_bytes()          # built BEFORE openpyxl is blocked
    monkeypatch.setattr(builtins, "__import__", fake)
    # A real workbook, so this exercises the zip/XML fallback rather than the
    # "not a zip" error path.
    res = P.parse("a.xlsx", raw)
    assert any("openpyxl" in w for w in res.warnings), res.warnings
    assert res.blocks, "the fallback must still read the sheet"
    t = res.blocks[0]
    assert t.header == ["col_a", "col_b"]
    assert ["1", "2"] in t.rows


def test_corrupt_xlsx_reports_instead_of_raising():
    res = P.parse("a.xlsx", b"definitely not a zip")
    assert not res.blocks
    assert res.warnings


# ── cross-cutting ────────────────────────────────────────────────────────────

@pytest.mark.parametrize("name, data", [
    ("a.md", MD.encode()),
    ("a.txt", b"Some heading\n\nA sentence here. Another one.\n"),
    ("a.html", HTML.encode()),
    ("a.csv", b"a,b\n1,2\n3,4\n"),
])
def test_every_parser_output_chunks_cleanly(name, data):
    res = P.parse(name, data)
    chunks = chunk_blocks(res.blocks, target_words=40)
    assert chunks
    for c in chunks:
        assert c.text.strip()
        assert len(c.text) < P.MAX_CHUNK_CHARS * 2


def test_parse_reports_its_type():
    assert P.parse("x.md", MD.encode()).doc_type == "md"
    assert P.parse("x.csv", b"a\n1\n").doc_type == "csv"


def test_bom_and_windows_line_endings_do_not_break_anything():
    for src in (MD, "Heading\n\nBody text. More body.\n"):
        for enc in ("utf-8-sig", "cp1252"):
            data = src.replace("\n", "\r\n").encode(enc, "replace")
            name = "a.md" if src is MD else "a.txt"
            res = P.parse(name, data)
            assert res.blocks, (enc, name)
